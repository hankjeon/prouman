import importlib.util
import os
import sys
import tempfile
import unittest
from pathlib import Path

import yaml
from openpyxl import load_workbook

# 창을 띄우지 않고 검증한다. 사용자가 지정한 플랫폼이 있으면 그대로 쓴다.
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
# PySide6가 없을 때만 건너뛴다. 설치됐는데 불러오지 못하면(예: libEGL 누락) 성공처럼 보이지 않게 오류를 낸다.
if importlib.util.find_spec('PySide6') is None:
    QApplication = None
else:
    from PySide6.QtCore import QItemSelectionModel, QSettings, Qt
    from PySide6.QtGui import QGuiApplication
    from PySide6.QtWidgets import QApplication, QComboBox

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
sys.path.insert(0, str(ROOT / 'tests'))

from test_allocation import person, project

from allocation.document import ProjectEntry
from allocation.loading import load_inputs

if QApplication is not None:
    from allocation import theme
    from allocation.gui import (DEFAULT_FONT_SIZE, MAX_FONT_SIZE, MIN_FONT_SIZE, MainWindow, SelectionDelegate,
                                find_projects_file, severity_colors)


@unittest.skipUnless(QApplication, 'PySide6가 설치되지 않아 GUI 테스트를 건너뜁니다.')
class GuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.path = Path(directory.name)
        self.settings = QSettings(str(self.path / 'settings.ini'), QSettings.IniFormat)
        self.warnings = []
        self.save_answers = []  # 저장하지 않은 변경 질문에 줄 답. 비면 'discard'
        self.confirm_answer = True
        self.save_paths = []  # 저장 위치 대화상자에 줄 경로

    def window(self):
        window = MainWindow(self.settings)
        window.warn = lambda title, text: self.warnings.append((title, text))
        window.ask_save = lambda kind: self.save_answers.pop(0) if self.save_answers else 'discard'
        window.confirm = lambda title, text: self.confirm_answer
        window.ask_save_path = lambda kind: self.save_paths.pop(0) if self.save_paths else ''
        self.addCleanup(window.deleteLater)
        return window

    def write(self, projects, people, year=2027):
        p, h = self.path / 'projects.yaml', self.path / 'people' / 'people.yaml'
        h.parent.mkdir(exist_ok=True)
        p.write_text(yaml.safe_dump({'projects': projects}, allow_unicode=True, sort_keys=False), encoding='utf-8')
        h.write_text(yaml.safe_dump({'year': year, 'people': people}, allow_unicode=True, sort_keys=False),
                     encoding='utf-8')
        return p, h

    def generated(self, projects, people):
        p, h = self.write(projects, people)
        window = self.window()
        window.choose_projects(str(p))
        window.choose_people(str(h))
        window.generate()
        return window

    @staticmethod
    def texts(view):
        model = view.model()
        return [[model.text(r, c) for c in range(model.columnCount())] for r in range(model.rowCount())]

    def test_T29_generate_views_copy_excel_and_settings(self):
        w = self.generated([project(), project('사업2')],
                           [person({'사업1': '0.60000', '사업2': '0.40000'}), person({'사업2': '1'}, name='김철수')])
        self.assertEqual([w.input_tabs.tabText(i) for i in range(w.input_tabs.count())], ['사업 입력', '사람 입력', '종합'])
        self.assertEqual([w.result_tabs.tabText(i) for i in range(w.result_tabs.count())],
                         ['검증', '요약', '개인상세', '사업상세'])
        self.assertEqual(w.pages.currentIndex(), w.RESULT_PAGE)
        self.assertEqual(w.result_tabs.currentIndex(), w.SUMMARY_TAB)
        self.assertTrue(w.excel_button.isEnabled())
        self.assertRegex(w.result_header.text(),
                         r'^관리 연도 2027 · ERROR 0건 / WARNING 0건 · \d\d:\d\d:\d\d 생성 · projects.yaml / people.yaml$')
        self.assertEqual(self.texts(w.issue_view), [])
        self.assertEqual(sorted(self.texts(w.people_summary_view)),
                         [['김철수', '1.00000', '0.00000', '1.00000'], ['홍길동', '1.00000', '0.00000', '1.00000']])
        self.assertEqual(self.texts(w.project_summary_view),
                         [['사업1', '사업1', '10.00000', '0.60000', '9.40000', ''],
                          ['사업2', '사업2', '10.00000', '1.40000', '8.60000', '']])
        projects = self.texts(w.project_view)
        self.assertEqual([len(row) for row in projects], [18, 18])
        self.assertEqual(projects[0][6:9], ['0.05000', '0.05000', '0.05000'])

        # 사람을 고르면 사업·공백·합계 표가 나오고 합계 행은 굵게 표시한다.
        self.assertEqual([w.person_list.item(i).text() for i in range(w.person_list.count())], ['김철수', '홍길동'])
        self.assertEqual(w.person_list.currentRow(), 0)
        self.assertEqual([row[0] for row in self.texts(w.person_view)], ['사업2', '공백', '합계'])
        w.person_list.setCurrentRow(1)
        detail = self.texts(w.person_view)
        self.assertEqual([row[0] for row in detail], ['사업1', '사업2', '공백', '합계'])
        self.assertEqual(detail[0][1:4], ['0.05000', '0.05000', '0.05000'])
        self.assertEqual(detail[3][-1], '1.00000')
        model = w.person_view.model()
        self.assertTrue(model.data(model.index(3, 1), Qt.FontRole).bold())
        self.assertIsNone(model.data(model.index(0, 1), Qt.FontRole))
        self.assertEqual(model.data(model.index(0, 1), Qt.TextAlignmentRole), int(Qt.AlignRight | Qt.AlignVCenter))
        self.assertIsNone(model.data(model.index(0, 0), Qt.TextAlignmentRole))

        # 선택한 범위를 탭 구분 텍스트로 복사하고, 선택하지 않은 칸은 비운다.
        view = w.person_view
        self.assertIsInstance(view.itemDelegate(), SelectionDelegate)
        select = QItemSelectionModel.SelectionFlag.Select
        view.selectionModel().select(view.model().index(0, 1), select)
        view.selectionModel().select(view.model().index(1, 2), select)
        self.assertEqual(view.copy_selection(), '0.05000\t\n\t0.03333')
        self.assertEqual(QGuiApplication.clipboard().text(), '0.05000\t\n\t0.03333')
        view.clearSelection()
        self.assertEqual(view.copy_selection(), '')

        self.assertTrue(w.save_excel(str(self.path / 'out.xlsx')))
        self.assertEqual(load_workbook(self.path / 'out.xlsx').sheetnames, ['검증', '요약', '개인상세', '사업상세'])
        (self.path / 'folder.xlsx').mkdir()
        self.assertFalse(w.save_excel(str(self.path / 'folder.xlsx')))
        self.assertEqual(self.warnings[-1][0], '엑셀 저장 실패')

        # 연 파일 경로를 기억해 다음 창에서 다시 연다.
        again = self.window()
        self.assertEqual(again.projects_edit.text(), str(self.path / 'projects.yaml'))
        self.assertEqual(again.people_edit.text(), str(self.path / 'people' / 'people.yaml'))
        self.assertEqual([p.alias for p in again.project_editor.projects], ['사업1', '사업2'])
        self.assertEqual(len(again.people_editor.doc.people), 2)

    def test_T29_errors_clear_results(self):
        w = self.generated([project()], [person()])
        self.assertTrue(w.excel_button.isEnabled())
        _, h = self.write([project()], [person({'사업1': '0.90000'})])
        w.choose_people(str(h))
        w.generate()
        self.assertEqual((w.pages.currentIndex(), w.result_tabs.currentIndex()), (w.RESULT_PAGE, w.ISSUES_TAB))
        self.assertFalse(w.excel_button.isEnabled())
        self.assertIsNone(w.result)
        self.assertIn('오류로 결과를 만들지 않았습니다', w.result_header.text())
        issues = self.texts(w.issue_view)
        self.assertIn(['ERROR', 'E101'], [row[:2] for row in issues])
        model = w.issue_view.model()
        self.assertEqual(model.data(model.index(0, 0), Qt.ForegroundRole), severity_colors()['ERROR'])
        for view in (w.people_summary_view, w.project_summary_view, w.person_view, w.project_view):
            self.assertEqual(view.model().rowCount(), 0)
        self.assertEqual(w.person_list.count(), 0)

        # 사업 입력이 비면 YAML 검증 오류(E003)만 나오고, 불완전한 사업 목록으로 미등록 약칭(E007)을 덧붙이지 않는다.
        w.new_document('projects')
        w.generate()
        codes = [row[1] for row in self.texts(w.issue_view)]
        self.assertIn('E003', codes)
        self.assertNotIn('E007', codes)
        self.assertFalse(w.excel_button.isEnabled())

        w.choose_projects(str(self.path / '없는 파일.yaml'))
        self.assertEqual(self.warnings[-1][0], '사업 파일 열기 실패')

    def test_T29_warning_only_shows_results(self):
        w = self.generated([project(excluded=[4, 4]), project('사업2')],
                           [person({'사업1': '0.60000', '사업2': '0.40000'})])
        self.assertTrue(w.excel_button.isEnabled())
        self.assertEqual([row[:2] for row in self.texts(w.issue_view)], [['WARNING', 'W001']])
        model = w.issue_view.model()
        self.assertEqual(model.data(model.index(0, 0), Qt.ForegroundRole), severity_colors()['WARNING'])
        self.assertEqual(w.people_summary_view.model().rowCount(), 1)

    def test_T29_font_size(self):
        w = self.generated([project()], [person()])
        self.assertEqual(DEFAULT_FONT_SIZE, 10)
        self.assertEqual(w.font().pointSize(), DEFAULT_FONT_SIZE)
        # 크기를 바꾸기 전에는 설정에 저장하지 않는다.
        self.assertFalse(self.settings.contains('font_size'))
        self.assertEqual(w.person_view.font().pointSize(), DEFAULT_FONT_SIZE)
        height = w.person_view.verticalHeader().defaultSectionSize()
        width = w.person_view.columnWidth(1)
        w.change_font_size(1)
        w.change_font_size(1)
        self.assertEqual(w.person_view.font().pointSize(), DEFAULT_FONT_SIZE + 2)
        # 행 높이와 열 너비가 새 글자 크기에 맞게 커진다.
        self.assertGreater(w.person_view.verticalHeader().defaultSectionSize(), height)
        self.assertGreater(w.person_view.columnWidth(1), width)
        # 합계 행의 굵은 글꼴도 창 글자 크기를 따른다.
        model = w.person_view.model()
        bold = model.data(model.index(model.rowCount() - 1, 1), Qt.FontRole).resolve(w.person_view.font())
        self.assertEqual((bold.pointSize(), bold.bold()), (DEFAULT_FONT_SIZE + 2, True))
        # 바꾼 크기는 다음 창에 이어지고, 범위를 벗어나지 않으며, Ctrl+0으로 기본값이 된다.
        self.assertEqual(self.window().font().pointSize(), DEFAULT_FONT_SIZE + 2)
        for _ in range(30):
            w.change_font_size(1)
        self.assertEqual(w.font_size, MAX_FONT_SIZE)
        for _ in range(30):
            w.change_font_size(-1)
        self.assertEqual(w.font_size, MIN_FONT_SIZE)
        w.change_font_size(0)
        self.assertEqual(w.font().pointSize(), DEFAULT_FONT_SIZE)
        menus = {a.text(): a.menu() for a in w.menuBar().actions()}
        self.assertEqual([a.text() for a in menus['보기'].actions()],
                         ['입력 화면', '결과 화면', '', '글자 크게', '글자 작게', '기본 크기'])

    # --- T30 입력 편집기 ---

    def matrix_texts(self, w):
        # 종합 탭(읽기 전용 사람 × 사업 표)
        m = w.people_editor.overview.matrix
        return [[m.item(r, c).text() for c in range(m.columnCount())] for r in range(m.rowCount())]

    def alloc_texts(self, w):
        # 사람 입력 탭에서 고른 사람의 참여 사업 표: 사업, 참여율, 사업명
        a = w.people_editor.allocations
        return [[a.item(r, c).text() for c in range(a.columnCount())] for r in range(a.rowCount())]

    def test_T30_edit_tables_and_generate_unsaved(self):
        p, h = self.write([project(), project('사업2')], [person({'사업1': '0.60000', '사업2': '0.40000'})])
        w = self.window()
        w.choose_people(str(h))
        editor, alloc, matrix = w.people_editor, w.people_editor.allocations, w.people_editor.overview.matrix
        self.assertEqual([w.project_editor.table.item(0, c).text() for c in range(4)],
                         ['사업1', '사업1', '2027-01-01', '2027-12-31'])
        self.assertEqual([matrix.horizontalHeaderItem(c).text() for c in range(matrix.columnCount())],
                         ['이름', '사업1', '사업2', '사업 합계', '공백', '합계'])
        self.assertEqual(self.matrix_texts(w), [['홍길동', '0.60000', '0.40000', '1.00000', '0.00000', '1.00000']])
        self.assertEqual(self.alloc_texts(w), [['사업1', '0.60000', '사업1'], ['사업2', '0.40000', '사업2']])
        self.assertEqual(editor.year_edit.text(), '2027')

        # 참여율을 고치면 합계가 바로 바뀌고 1이 아니면 표시한다. 저장 전이어도 생성은 화면 내용으로 계산한다.
        alloc.item(0, 1).setText('0.5')
        self.assertEqual(editor.doc.people[0].allocations['사업1'], '0.5')
        self.assertEqual(self.matrix_texts(w)[0][3:], ['0.90000', '0.00000', '0.90000'])
        self.assertTrue(matrix.item(0, 5).toolTip())
        self.assertIn('⚠', editor.totals.text())
        self.assertTrue(editor.people_list.item(0).toolTip())
        self.assertEqual(w.input_tabs.tabText(w.PEOPLE_TAB), '사람 입력 *')
        w.generate()
        self.assertIn('E101', [row[1] for row in self.texts(w.issue_view)])
        alloc.item(1, 1).setText('0.5')
        self.assertEqual(matrix.item(0, 5).toolTip(), '')
        self.assertNotIn('⚠', editor.totals.text())
        self.assertEqual(editor.people_list.item(0).toolTip(), '')
        w.generate()
        self.assertTrue(w.excel_button.isEnabled())
        self.assertIn('0.60000', h.read_text(encoding='utf-8'))  # 저장하기 전에는 파일이 그대로다.

        # 형식이 틀린 값은 칸을 표시하고 안내를 띄운다. 빈 참여율도 안내하고, 행 삭제로 사업을 뺀다.
        alloc.item(0, 1).setText('0.5.1')
        self.assertTrue(alloc.item(0, 1).toolTip())
        self.assertEqual(self.matrix_texts(w)[0][3], '?')
        alloc.item(0, 1).setText('')
        self.assertIn('참여율을 입력하세요', alloc.item(0, 1).toolTip())
        alloc.setCurrentCell(0, 0)
        editor.remove_row()
        self.assertEqual(editor.doc.people[0].allocations, {'사업2': '0.5'})

        # 행 추가 → 사업 고르기(드롭다운은 이 사람이 이미 고른 사업을 빼고 잔여 배정을 보인다) → 붙여넣기.
        editor.add_row()
        self.assertEqual(list(editor.doc.people[0].allocations), ['사업2', ''])
        self.assertIn('사업을 고르세요', alloc.item(1, 0).toolTip())
        self.assertEqual(editor.choices(1), [('사업1', '사업1 — 사업1 (잔여 10.00000)')])
        alloc.item(1, 0).setText('사업1')
        self.assertEqual(list(editor.doc.people[0].allocations), ['사업2', '사업1'])
        alloc.item(1, 0).setText('사업2')  # 같은 사업은 한 사람에게 두 번 넣을 수 없다.
        self.assertIn('이미 있습니다', editor.message.text())
        self.assertEqual(list(editor.doc.people[0].allocations), ['사업2', '사업1'])
        alloc.setCurrentCell(0, 1)
        alloc.paste('0.25\n0.75\n9')
        self.assertEqual(self.alloc_texts(w), [['사업2', '0.25', '사업2'], ['사업1', '0.75', '사업1']])
        self.assertEqual(self.matrix_texts(w)[0][1:6], ['0.75', '0.25', '1.00000', '0.00000', '1.00000'])
        alloc.clearSelection()
        alloc.selectRow(0)
        self.assertEqual(alloc.copy_selection(), '사업2\t0.25\t사업2')
        # 종합 탭은 읽기 전용이고, 행을 두 번 누르면 그 사람의 입력 화면으로 간다.
        self.assertFalse(matrix.item(0, 1).flags() & Qt.ItemIsEditable)
        w.input_tabs.setCurrentIndex(w.OVERVIEW_TAB)
        matrix.cellDoubleClicked.emit(0, 1)
        self.assertEqual(w.input_tabs.currentIndex(), w.PEOPLE_TAB)

    def test_T30_absences_and_year(self):
        _, h = self.write([project(year=2026)], [person({'사업1': '0.95430'})], year=2026)
        w = self.window()
        w.choose_people(str(h))
        editor = w.people_editor
        self.assertEqual(self.matrix_texts(w)[0][2:], ['0.95430', '0.00000', '0.95430'])
        editor.add_absence()
        absences = editor.absences
        absences.item(0, 0).setText('2026-07-15')
        self.assertTrue(absences.item(0, 1).toolTip())  # 종료일이 비어 있음
        absences.item(0, 1).setText('2026.7.31')
        absences.item(0, 2).setText('출산휴가')
        self.assertEqual(self.matrix_texts(w)[0][2:], ['0.95430', '0.04570', '1.00000'])
        self.assertEqual(editor.doc.people[0].absences[0].reason, '출산휴가')
        editor.year_edit.setText('20x6')
        editor.year_edited('20x6')
        self.assertEqual(self.matrix_texts(w)[0][3], '?')
        self.assertTrue(editor.year_edit.toolTip())
        editor.year_edited('2026')
        w.generate()
        self.assertTrue(w.excel_button.isEnabled())

    def test_T30_project_editing_rename_and_remove(self):
        _, h = self.write([project(), project('사업2')], [person({'사업1': '0.60000', '사업2': '0.40000'})])
        w = self.window()
        w.choose_people(str(h))
        table, grid = w.project_editor.table, w.project_editor.grid
        # 약칭을 바꾸면 사람 참여율도 따라간다.
        table.item(0, 1).setText('연구')
        self.assertEqual(w.people_editor.doc.people[0].allocations, {'연구': '0.60000', '사업2': '0.40000'})
        self.assertEqual(w.input_tabs.tabText(w.PROJECTS_TAB), '사업 입력 *')
        self.assertEqual(w.input_tabs.tabText(w.PEOPLE_TAB), '사람 입력 *')
        # 입력 확인: 약칭 공백·중복, 날짜, 연도·배정·제외 월.
        # 다른 사업과 같은 약칭을 거쳐도 그 사업의 참여율이 옮겨지지 않는다.
        table.item(1, 1).setText('연구')
        self.assertTrue(table.item(1, 1).toolTip())
        table.item(1, 1).setText('사업 2')
        self.assertTrue(table.item(1, 1).toolTip())
        table.item(1, 1).setText('사업2')
        self.assertEqual(w.people_editor.doc.people[0].allocations, {'연구': '0.60000', '사업2': '0.40000'})
        table.item(1, 3).setText('2026-12-31')
        self.assertIn('시작일', table.item(1, 3).toolTip())
        table.item(1, 3).setText('2027-12-31')
        # 연도 표: 관리 연도 2027이면 2026~2030, 제외월 열은 2027.
        grid.setCurrentCell(1, 1)
        grid.item(1, 5).setText('4, 13')
        self.assertTrue(grid.item(1, 5).toolTip())
        grid.item(1, 5).setText('4, 5, 6')
        grid.item(1, 2).setText('1')
        self.assertEqual([e.year for e in w.project_editor.projects[1].years], ['2027', '2028'])
        grid.item(1, 2).setText('')
        self.assertEqual([(e.year, e.excluded) for e in w.project_editor.projects[1].years], [('2027', '4, 5, 6')])
        # 사업2를 4~6월 제외하면 사업1(연구)이 그 달을 채운다.
        w.generate()
        self.assertTrue(w.excel_button.isEnabled(), self.texts(w.issue_view))
        # 참여율이 있는 사업을 지울 때는 확인하고, 취소하면 그대로 둔다.
        self.confirm_answer = False
        w.project_editor.remove_project()
        self.assertEqual(len(w.project_editor.projects), 2)
        self.confirm_answer = True
        w.project_editor.remove_project()
        self.assertEqual([p.alias for p in w.project_editor.projects], ['연구'])
        self.assertEqual(w.people_editor.doc.people[0].allocations, {'연구': '0.60000'})

    def test_T30_unknown_alias_and_dropped_fields(self):
        p, h = self.write([project()], [person({'사업1': '0.5', '없는사업': '0.5'})])
        text = p.read_text(encoding='utf-8').replace('- name: 사업1', '- id: 7\n  name: 사업1')
        p.write_text(text, encoding='utf-8')
        w = self.window()
        w.choose_people(str(h))
        self.assertEqual(self.warnings[-1][0], '지원하지 않는 항목')
        self.assertIn('projects[0].id', self.warnings[-1][1])
        matrix = w.people_editor.overview.matrix
        headers = [matrix.horizontalHeaderItem(c).text() for c in range(3)]
        self.assertEqual(headers, ['이름', '사업1', '없는사업 (미등록)'])
        self.assertTrue(matrix.item(0, 2).toolTip())
        self.assertIn('사업 목록에 없는 약칭', w.people_editor.allocations.item(1, 0).toolTip())
        w.generate()
        self.assertIn('E007', [row[1] for row in self.texts(w.issue_view)])
        # 그 행을 지우면 미등록 열이 사라진다.
        w.people_editor.allocations.setCurrentCell(1, 0)
        w.people_editor.remove_row()
        self.assertEqual(matrix.columnCount(), 5)

    def test_T30_save_backup_and_new_documents(self):
        p, h = self.write([project()], [person()])
        original = p.read_text(encoding='utf-8')
        w = self.window()
        w.choose_people(str(h))
        w.project_editor.table.item(0, 0).setText('연구 지원 사업')
        self.assertTrue(w.save('projects'))
        self.assertEqual((self.path / 'projects.yaml.bak').read_text(encoding='utf-8'), original)
        self.assertEqual(w.input_tabs.tabText(w.PROJECTS_TAB), '사업 입력')
        self.assertIn('"연구 지원 사업"', p.read_text(encoding='utf-8'))
        # 같은 실행에서 다시 저장해도 첫 원본 백업은 그대로 둔다.
        w.project_editor.table.item(0, 0).setText('운영')
        self.assertTrue(w.save_all())
        self.assertEqual((self.path / 'projects.yaml.bak').read_text(encoding='utf-8'), original)

        # 새로 만들기 → 입력 → 다른 이름으로 저장한 파일은 CLI와 같은 규칙으로 읽힌다.
        w.new_document('projects')
        w.new_document('people')
        self.assertEqual(w.projects_edit.text(), '')
        pe, me = w.project_editor, w.people_editor
        pe.add_project()
        for column, value in enumerate(['새 사업', '새사업', '2026-01-01', '2026-12-31']):
            pe.table.item(0, column).setText(value)
        pe.grid.item(0, 1).setText('2')  # 관리 연도 2026 → 2025~2029의 두 번째 열
        me.year_edited('2026')
        me.add_person()
        self.assertEqual(me.doc.people[0].name, '새 사람')
        me.people_list.item(0).setText('이영희')
        self.assertEqual(me.doc.people[0].name, '이영희')
        me.add_row()
        self.assertEqual([alias for alias, _ in me.choices(0)], ['새사업'])
        me.allocations.item(0, 0).setText('새사업')
        me.allocations.item(0, 1).setText('1')
        w.generate()
        self.assertTrue(w.excel_button.isEnabled(), self.texts(w.issue_view))
        folder = self.path / 'new'
        folder.mkdir()
        self.save_paths = [str(folder / 'projects.yaml'), str(folder / 'people_2026.yaml')]
        self.assertTrue(w.save_all())
        inputs, issues = load_inputs(folder / 'projects.yaml', folder / 'people_2026.yaml')
        self.assertIsNotNone(inputs, issues)
        self.assertEqual(inputs.people['이영희'].targets, {'새사업': 100000})
        self.assertFalse((folder / 'projects.yaml.bak').exists())

        # 저장하지 않은 변경이 있으면 다른 파일을 열기 전에 묻는다.
        me.allocations.item(0, 1).setText('0.9')
        self.save_answers = ['cancel']
        w.choose_people(str(h))
        self.assertEqual(w.people_edit.text(), str(folder / 'people_2026.yaml'))
        self.save_answers = ['discard']
        w.choose_people(str(h))
        self.assertEqual(w.people_edit.text(), str(h))

    # --- T31 입력 화면과 결과 화면 ---

    def test_T31_pages_and_stale_results(self):
        _, h = self.write([project()], [person()])
        w = self.window()
        self.assertEqual(w.pages.currentIndex(), w.INPUT_PAGE)
        self.assertTrue(w.nav_buttons[w.INPUT_PAGE].isChecked())
        # 생성 전 결과 화면에는 안내 문구만 있다.
        w.show_page(w.RESULT_PAGE)
        self.assertTrue(w.nav_buttons[w.RESULT_PAGE].isChecked())
        self.assertEqual(w.result_stack.currentIndex(), 0)
        self.assertEqual(w.result_header.text(), '결과 없음')
        self.assertFalse(w.excel_button.isEnabled())
        w.show_page(w.INPUT_PAGE)

        w.choose_people(str(h))
        self.assertTrue(w.stale_bar.isHidden())  # 생성 전에는 안내하지 않는다.
        w.generate()
        self.assertEqual((w.pages.currentIndex(), w.result_stack.currentIndex()), (w.RESULT_PAGE, 1))
        self.assertTrue(w.stale_bar.isHidden())
        # 생성 뒤 입력을 고치면 결과가 예전 입력 기준임을 알린다. 메뉴에도 저장 안 한 표시가 붙는다.
        w.people_editor.allocations.item(0, 1).setText('0.5')
        self.assertFalse(w.stale_bar.isHidden())
        self.assertIn('생성한 입력 기준', w.stale_label.text())
        self.assertEqual(w.nav_buttons[w.INPUT_PAGE].text(), '✎  입력 *')
        self.assertTrue(w.excel_button.isEnabled())  # 이전 결과는 그대로 저장할 수 있다.
        # 다시 생성하면 안내가 사라지고, 오류면 검증 탭과 안내 문구를 보여 준다.
        w.generate()
        self.assertTrue(w.stale_bar.isHidden())
        self.assertEqual(w.result_tabs.currentIndex(), w.ISSUES_TAB)
        self.assertIn('입력 화면에서 고친 뒤 다시 생성하세요', w.result_header.text())
        self.assertFalse(w.excel_button.isEnabled())
        # 저장하면 메뉴 표시가 사라지고, 결과는 그대로라 안내도 생기지 않는다.
        w.people_editor.allocations.item(0, 1).setText('1')
        w.generate()
        self.assertTrue(w.save('people'))
        self.assertEqual(w.nav_buttons[w.INPUT_PAGE].text(), '✎  입력')
        self.assertTrue(w.stale_bar.isHidden())
        # 다른 파일을 열거나 새로 만들어도 안내한다. 새로 만들기는 입력 화면으로 돌아간다.
        w.choose_people(str(h))
        self.assertFalse(w.stale_bar.isHidden())
        w.generate()
        w.new_document('people')
        self.assertFalse(w.stale_bar.isHidden())
        self.assertEqual((w.pages.currentIndex(), w.input_tabs.currentIndex()), (w.INPUT_PAGE, w.PEOPLE_TAB))
        # 새 파일로 생성하면 머리 줄에 새 파일로 표시한다.
        w.generate()
        self.assertIn('(새 사람 파일)', w.result_header.text())
        menus = {a.text(): a.menu() for a in w.menuBar().actions()}
        actions = {a.text(): a for a in menus['보기'].actions()}
        self.assertEqual(actions['입력 화면'].shortcut().toString(), 'Ctrl+1')
        actions['입력 화면'].trigger()
        self.assertEqual(w.pages.currentIndex(), w.INPUT_PAGE)
        actions['결과 화면'].trigger()
        self.assertEqual(w.pages.currentIndex(), w.RESULT_PAGE)

    # --- T32 사업 입력의 연도별 배정인력 표 ---

    def grid_texts(self, w):
        g = w.project_editor.grid
        return [[g.item(r, c).text() for c in range(g.columnCount())] for r in range(g.rowCount())]

    def test_T32_year_grid(self):
        projects = [project(period='2026-04-15~2028-09-10'), project('사업2', period='2025-01-01~2029-12-31')]
        projects[0]['yearly'] = {2026: {'allocated_manpower': '1', 'excluded_months': []},
                                 2027: {'allocated_manpower': '1.5', 'excluded_months': [4, 5, 6]},
                                 2032: {'allocated_manpower': '0.5'}}
        _, h = self.write(projects, [person({'사업1': '0.60000', '사업2': '0.40000'})])
        w = self.window()
        w.choose_people(str(h))
        pe = w.project_editor
        table, grid = pe.table, pe.grid
        headers = [grid.horizontalHeaderItem(c).text() for c in range(grid.columnCount())]
        self.assertEqual(headers, ['2026', '2027', '2028', '2029', '2030', '2027 제외월'])
        self.assertTrue(grid.horizontalHeaderItem(1).font().bold())  # 관리 연도
        self.assertEqual(self.grid_texts(w), [['1', '1.5', '', '', '', '4, 5, 6'],
                                              ['', '10.00000', '', '', '', '']])
        # 총 배정은 범위 밖(2032) 연도까지 더하고 편집할 수 없다.
        self.assertEqual(table.item(0, pe.TOTAL).text(), '3.00000')
        self.assertFalse(table.item(0, pe.TOTAL).flags() & Qt.ItemIsEditable)
        # 범위 밖 값은 ▶ 버튼에 건수로 알린다. 한 해씩 옮기면 제외월 기준 연도는 보이는 범위 안에 남는다.
        self.assertEqual((pe.prev_button.text(), pe.next_button.text()), ('◀', '1 ▶'))
        self.assertEqual(pe.range_label.text(), '2026 ~ 2030')
        for _ in range(2):
            pe.shift_years(1)
        self.assertEqual(pe.range_label.text(), '2028 ~ 2032')
        self.assertEqual(pe.prev_button.text(), '◀ 3')  # 사업1 2026·2027, 사업2 2027
        self.assertEqual(pe.excluded_year, 2028)
        self.assertEqual(self.grid_texts(w)[0][4], '0.5')
        pe.reset_years()
        self.assertEqual(pe.range_label.text(), '2026 ~ 2030')

        # 색: 기간 밖(회색), 제외 월 있음(연보라, 기간 밖보다 우선), 오류(붉은색, 가장 우선).
        self.assertEqual(grid.item(0, 1).background().color(), theme.excluded_color())
        self.assertIn('제외 월: 4, 5, 6', grid.item(0, 1).toolTip())
        self.assertEqual(grid.item(0, 3).background().color(), theme.outside_color())
        self.assertIn('사업 기간 밖', grid.item(0, 3).toolTip())
        grid.item(0, 1).setText('1.5.5')
        self.assertEqual(grid.item(0, 1).background().color(), theme.problem_color())
        self.assertEqual(table.item(0, pe.TOTAL).text(), '?')
        self.assertIn('2027', table.item(0, pe.TOTAL).toolTip())
        grid.item(0, 1).setText('1.5')
        self.assertEqual(table.item(0, pe.TOTAL).text(), '3.00000')

        # 연도 열을 옮기면 제외월 열 제목·내용이 그 연도를 따른다. 항목이 없는 사업은 편집할 수 없다.
        grid.setCurrentCell(0, 0)
        self.assertEqual(grid.horizontalHeaderItem(5).text(), '2026 제외월')
        self.assertEqual(self.grid_texts(w)[0][5], '')
        self.assertFalse(grid.item(1, 5).flags() & Qt.ItemIsEditable)
        self.assertIn('먼저 2026 배정인력', grid.item(1, 5).toolTip())
        grid.item(0, 5).setText('3,4')
        self.assertEqual(grid.item(0, 0).background().color(), theme.excluded_color())
        grid.setCurrentCell(0, 5)  # 제외월 열로 가도 기준 연도는 유지한다.
        self.assertEqual(pe.excluded_year, 2026)
        grid.item(0, 5).setText('')
        self.assertNotEqual(grid.item(0, 0).background().color(), theme.excluded_color())

        # 배정 칸: 입력하면 항목 생성, 비우면 삭제, 제외 월이 남아 있으면 빈 배정으로 남아 오류 표시.
        grid.item(1, 2).setText('2')
        self.assertEqual([(e.year, e.manpower) for e in pe.projects[1].years], [('2027', '10.00000'), ('2028', '2')])
        self.assertEqual(table.item(1, pe.TOTAL).text(), '12.00000')
        grid.item(1, 2).setText('')
        self.assertEqual([e.year for e in pe.projects[1].years], ['2027'])
        grid.setCurrentCell(0, 1)
        grid.item(0, 1).setText('')
        self.assertEqual([(e.year, e.manpower, e.excluded) for e in pe.projects[0].years if e.year == '2027'],
                         [('2027', '', '4, 5, 6')])
        self.assertEqual(grid.item(0, 1).background().color(), theme.problem_color())
        grid.item(0, 1).setText('1.5')

        # 사업 기간을 바꾸면 기간 밖 표시가 바뀐다.
        table.item(0, pe.END).setText('2029-12-31')
        self.assertNotEqual(grid.item(0, 3).background().color(), theme.outside_color())

        # 좌우 현재 행과 세로 스크롤이 맞물린다.
        table.setCurrentCell(1, 0)
        self.assertEqual(grid.currentRow(), 1)
        grid.setCurrentCell(0, 2)
        self.assertEqual(table.currentRow(), 0)
        self.assertEqual(table.horizontalHeader().height(), grid.horizontalHeader().height())
        for n in range(40):
            pe.projects.append(ProjectEntry(f'추가 사업 {n}', f'추가{n}', '2025-01-01', '2029-12-31'))
        pe.refresh()
        # 스크롤 범위는 화면에 배치된 뒤에 정해진다.
        w.resize(1200, 600)
        w.show()
        self.app.processEvents()
        self.assertGreater(grid.verticalScrollBar().maximum(), 7)
        self.assertEqual(table.verticalScrollBar().maximum(), grid.verticalScrollBar().maximum())
        grid.verticalScrollBar().setValue(7)
        self.assertEqual(table.verticalScrollBar().value(), 7)
        table.verticalScrollBar().setValue(3)
        self.assertEqual(grid.verticalScrollBar().value(), 3)
        del pe.projects[2:]
        pe.refresh()

        # 엑셀 범위 붙여넣기로 여러 사업·연도를 입력한다.
        grid.setCurrentCell(0, 2)
        grid.paste('2\t2\n3\t3')
        self.assertEqual(self.grid_texts(w)[0][2:4], ['2', '2'])
        self.assertEqual(self.grid_texts(w)[1][2:4], ['3', '3'])
        self.assertEqual(table.item(1, pe.TOTAL).text(), '16.00000')

        # 제외 월 4~6이 계산에 반영된다(사업1 2027 연구 지원은 4~6월 0, 그 달은 사업2가 채운다).
        w.generate()
        self.assertTrue(w.excel_button.isEnabled(), self.texts(w.issue_view))
        self.assertEqual(w.result.allocations['홍길동']['사업1'][3:6], [0, 0, 0])

    def test_T32_unshown_year_keys(self):
        p, h = self.write([project()], [person()])
        text = p.read_text(encoding='utf-8').replace('2027:', "'20x7':\n      allocated_manpower: '1'\n    2027:")
        p.write_text(text, encoding='utf-8')
        w = self.window()
        w.choose_people(str(h))
        self.assertIn('표시할 수 없는 연도 항목 1건', w.project_editor.notice.text())
        # 표시할 수 없는 항목도 총 배정에는 더한다(10 + 1).
        self.assertEqual(w.project_editor.table.item(0, w.project_editor.TOTAL).text(), '11.00000')
        w.generate()
        self.assertIn('E003', [row[1] for row in self.texts(w.issue_view)])

    # --- T33 사람별 입력 ---

    def test_T33_people_list_and_project_choices(self):
        projects = [project('운영'), project('연구', period='2026-01-01~2026-12-31'), project('예산없음')]
        projects[2]['yearly'] = {2028: {'allocated_manpower': '1'}}
        _, h = self.write(projects, [person({'운영': '0.7', '연구': '0.3'}),
                                     person({'운영': '1'}, name='김철수')])
        w = self.window()
        w.choose_people(str(h))
        editor = w.people_editor
        self.assertEqual([editor.people_list.item(i).text() for i in range(editor.people_list.count())],
                         ['홍길동', '김철수'])
        self.assertEqual(editor.title.text(), '홍길동의 참여 사업')
        # 관리 연도(2027) 배정이 없거나 사업 기간이 그 해가 아닌 사업은 드롭다운에 나오지 않는다.
        self.assertEqual([alias for alias, _ in editor.choices(0)], ['운영'])
        self.assertIn('배정이 없거나 사업 기간이 아닌 사업', editor.allocations.item(1, 0).toolTip())
        self.assertTrue(editor.people_list.item(0).toolTip())  # 문제가 있는 사람은 목록에서 표시
        # 잔여 배정 = 배정 10 - (0.7 + 1)
        self.assertIn('잔여 8.30000', editor.choices(0)[0][1])
        # 드롭다운 편집기는 고를 수 있는 사업을 약칭 값으로 담는다.
        combo = editor.allocations.itemDelegateForColumn(editor.PROJECT).createEditor(
            editor.allocations.viewport(), None, editor.allocations.model().index(0, 0))
        self.assertIsInstance(combo, QComboBox)
        self.assertEqual([combo.itemData(i) for i in range(combo.count())], ['운영'])
        combo.deleteLater()

        # 이름: 비어 있거나 다른 사람과 같으면 받지 않고 원래 이름으로 돌린다.
        editor.people_list.item(1).setText('홍길동')
        self.assertEqual(editor.doc.people[1].name, '김철수')
        self.assertEqual(editor.people_list.item(1).text(), '김철수')
        self.assertIn('이미 있는 이름', editor.message.text())
        editor.people_list.item(1).setText('  ')
        self.assertEqual(editor.doc.people[1].name, '김철수')
        editor.people_list.item(1).setText('김영희')
        self.assertEqual(editor.doc.people[1].name, '김영희')
        self.assertEqual(editor.message.text(), '')
        # 새 사람은 겹치지 않는 이름으로 추가된다.
        editor.add_person()
        editor.add_person()
        self.assertEqual([p.name for p in editor.doc.people][2:], ['새 사람', '새 사람 2'])
        self.assertEqual(editor.people_list.currentRow(), 3)
        editor.remove_person()
        self.assertEqual(len(editor.doc.people), 3)
        self.assertEqual(w.people_editor.overview.matrix.rowCount(), 3)

        # 휴직은 고른 사람에게만 들어간다.
        editor.select_person(1)
        editor.add_absence()
        editor.absences.item(0, 0).setText('2027-01-01')
        editor.absences.item(0, 1).setText('2027-12-31')
        self.assertEqual(len(editor.doc.people[1].absences), 1)
        self.assertEqual(editor.doc.people[0].absences, [])
        self.assertIn('공백 1.00000', editor.totals.text())

    def test_T29_find_projects_file(self):
        p, h = self.write([project()], [person()])
        w = self.window()
        w.choose_people(str(h))
        self.assertEqual(w.projects_edit.text(), str(p.resolve()))
        # 자동으로 채운 값은 다른 사람 파일을 고를 때 다시 찾는다.
        (self.path / 'people' / 'projects.yaml').write_text('', encoding='utf-8')
        w.choose_people(str(h))
        self.assertEqual(w.projects_edit.text(), str((self.path / 'people' / 'projects.yaml').resolve()))
        # 직접 고른 사업 파일은 덮어쓰지 않는다.
        chosen = self.path / 'other' / 'projects.yaml'
        chosen.parent.mkdir()
        chosen.write_bytes(p.read_bytes())
        w.choose_projects(str(chosen))
        w.choose_people(str(h))
        self.assertEqual(w.projects_edit.text(), str(chosen))
        # 사람 파일 폴더와 바로 위 폴더만 찾고 그보다 위는 보지 않는다.
        empty = self.path / 'a' / 'b' / 'people.yaml'
        empty.parent.mkdir(parents=True)
        self.assertIsNone(find_projects_file(empty))


@unittest.skipUnless(QApplication, 'PySide6가 설치되지 않아 GUI 테스트를 건너뜁니다.')
class ThemeTests(unittest.TestCase):
    """시스템 테마 대응 회귀 테스트. 다크에서도 메뉴·선택·칸이 보여야 한다."""

    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def themed(self, dark, func, *args):
        original = theme.is_dark
        theme.is_dark = lambda: dark
        try:
            return func(*args)
        finally:
            theme.is_dark = original

    def test_T29_theme_pairs_differ(self):
        # 라이트/다크 짝이 다르면 테마에 따라 다른 색이 나온다.
        for func in (theme.current_row_color, theme.current_cell_color, theme.problem_color,
                     theme.computed_color, theme.excluded_color, theme.outside_color,
                     theme.error_text, theme.warning_text):
            with self.subTest(func=func.__name__):
                self.assertNotEqual(self.themed(False, func).name(), self.themed(True, func).name())

    def test_T29_selection_follows_system_highlight(self):
        # 선택 배경은 시스템 강조색이라 어떤 테마에서도 어울리고,
        # 글자는 강조 글자색이라 배경과 대비된다.
        from PySide6.QtGui import QPalette
        palette = self.app.palette()
        self.assertEqual(theme.selection_color().name(), palette.color(QPalette.Highlight).name())
        self.assertEqual(theme.selected_text_color().name(), palette.color(QPalette.HighlightedText).name())
        self.assertNotEqual(theme.selection_color().name(), theme.selected_text_color().name())

    def test_T29_nav_stylesheet_uses_palette(self):
        # 좌측 메뉴: 배경·글자·선택이 시스템 팔레트에서 나온다.
        # 예전처럼 밝은 배경에 흰 글자(안 보임)가 되지 않는다.
        from PySide6.QtGui import QPalette
        sheet = theme.nav_stylesheet()
        palette = self.app.palette()
        for role in (QPalette.Window, QPalette.WindowText, QPalette.Highlight, QPalette.HighlightedText):
            self.assertIn(palette.color(role).name(), sheet)
        checked_bg = palette.color(QPalette.Highlight).name()
        checked_text = palette.color(QPalette.HighlightedText).name()
        self.assertIn(f'QToolButton:checked {{ background-color: {checked_bg}; color: {checked_text};', sheet)

    def test_T29_legend_text_contrasts(self):
        # 범례 칩은 글자색을 함께 지정한다. 밝은 칩에 검은 글자, 어두운 칩에 흰 글자.
        from PySide6.QtGui import QColor
        self.assertIn('color: #000000', theme.legend_stylesheet(QColor('#e4e4e4')))
        self.assertIn('color: #ffffff', theme.legend_stylesheet(QColor('#333333')))

    def test_T29_dark_error_text_is_light(self):
        # 다크에서 오류 빨강(#c62828)은 안 보이므로 밝은 빨강을 쓴다.
        dark_error = self.themed(True, theme.error_text)
        self.assertGreater(dark_error.lightness(), 100)


if __name__ == '__main__':
    unittest.main()
