import copy
import io
import itertools
import os
import random
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import yaml
from openpyxl import load_workbook

# src 레이아웃을 설치 없이 테스트합니다. CLI는 src/main.py에서 실행합니다.
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))

from allocation.cli import main
from allocation.engine import BASE, calculate, validate_result
from allocation.excel import write_workbook
from allocation.flow import transport
from allocation.loading import load_inputs
from allocation.models import amount
from allocation.rendering import issue_sort_key, person_summary, person_table, project_summary, render_result


def project(alias='사업1', year=2027, excluded=(), limit='10.00000', period=None):
    return dict(name=alias, alias=alias, period=period or f'{year}-01-01~{year}-12-31',
                yearly={year: dict(allocated_manpower=limit, excluded_months=list(excluded))})


def person(targets=None, name='홍길동', absences=()):
    return dict(name=name, allocations=targets if targets is not None else {'사업1': '1.00000'},
                absences=[dict(period=p) for p in absences])


class AcceptanceTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name)

    def write(self, projects, people, year=2027):
        p, h = self.path / 'projects.yaml', self.path / 'people.yaml'
        p.write_text(yaml.safe_dump({'projects': projects}, allow_unicode=True, sort_keys=False), encoding='utf-8')
        h.write_text(yaml.safe_dump({'year': year, 'people': people}, allow_unicode=True, sort_keys=False), encoding='utf-8')
        return p, h

    def run_files(self, paths):
        inputs, issues = load_inputs(*paths)
        return calculate(inputs, issues) if inputs else None, issues

    def ok(self, projects=None, people=None, year=2027):
        projects = projects if projects is not None else [project(year=year)]
        people = people if people is not None else [person()]
        paths = self.write(projects, people, year)
        result, issues = self.run_files(paths)
        self.assertIsNotNone(result, issues)
        self.assertFalse(result.has_errors, result.issues)
        self.assertEqual(validate_result(result), [])
        # 모든 성공 fixture에 대해 반복 및 YAML 순서 변경의 결정성을 확인한다.
        repeated, _ = self.run_files(paths)
        self.assertEqual(result.allocations, repeated.allocations)
        shuffled_projects = copy.deepcopy(list(reversed(projects)))
        shuffled_people = copy.deepcopy(list(reversed(people)))
        for p in shuffled_people:
            p['allocations'] = dict(reversed(list(p['allocations'].items())))
            if 'absences' in p:
                p['absences'].reverse()
        shuffled, _ = self.run_files(self.write(shuffled_projects, shuffled_people, year))
        self.assertIsNotNone(shuffled)
        self.assertFalse(shuffled.has_errors, shuffled.issues)
        self.assertEqual(result.allocations, shuffled.allocations)
        self.assertEqual(result.absence, shuffled.absence)
        # 출력량과 모든 합계는 5자리 문자열로도 검사한다.
        for name, rows in result.allocations.items():
            for m in range(12):
                total = result.absence[name][m] + sum(row[m] for row in rows.values())
                self.assertEqual(amount(total), '0.08334' if m % 3 == 2 else '0.08333')
            for alias, row in rows.items():
                target = result.inputs.people[name].targets[alias]
                self.assertEqual(amount(sum(row)), amount(target))
        return result

    def bad(self, projects=None, people=None, year=2027, code=None):
        projects = projects if projects is not None else [project()]
        people = people if people is not None else [person()]
        result, issues = self.run_files(self.write(projects, people, year))
        issues = result.issues if result else issues
        self.assertTrue(any(i.severity == 'ERROR' for i in issues), issues)
        if code:
            self.assertIn(code, [i.code for i in issues])
        if result:
            self.assertEqual(result.allocations, {})
            self.assertNotIn('개인별 연간 요약', render_result(result))
        return issues

    def test_T01_single_project(self):
        result = self.ok()
        self.assertEqual(result.allocations['홍길동']['사업1'], list(BASE))

    def test_T02_uniform(self):
        r = self.ok([project(), project('사업2')], [person({'사업1': '0.60000', '사업2': '0.40000'})])
        self.assertEqual(r.allocations['홍길동']['사업1'], [5000] * 12)
        self.assertEqual(r.allocations['홍길동']['사업2'], [b - 5000 for b in BASE])

    def test_T03_excluded_quarter(self):
        r = self.ok([project(excluded=[4, 5, 6]), project('사업2')],
                    [person({'사업1': '0.60000', '사업2': '0.40000'})])
        self.assertEqual(r.allocations['홍길동']['사업1'], [6667] * 3 + [0] * 3 + [6667] * 3 + [6666] * 3)

    def test_T04_partial_months(self):
        r = self.ok([project(period='2027-04-15~2027-09-10'), project('사업2')],
                    [person({'사업1': '0.50000', '사업2': '0.50000'})])
        self.assertEqual(r.allocations['홍길동']['사업1'], [0] * 3 + list(BASE[3:9]) + [0] * 3)

    def test_T05_partial_absence(self):
        r = self.ok(people=[person({'사업1': '0.95430'}, absences=['2026-07-15~2026-07-31'])],
                    projects=[project(year=2026)], year=2026)
        self.assertEqual(amount(r.absence['홍길동'][6]), '0.04570')
        self.assertEqual(amount(r.allocations['홍길동']['사업1'][6]), '0.03763')

    def test_T06_leap_year(self):
        r = self.ok([project(year=2028)], [person({'사업1': '0.99713'}, absences=['2028-02-29~2028-02-29'])], 2028)
        self.assertEqual(amount(r.absence['홍길동'][1]), '0.00287')

    def test_T07_overlapping_absences(self):
        r = self.ok(people=[person({'사업1': '0.95968'}, absences=['2027-07-01~2027-07-10', '2027-07-05~2027-07-15'])])
        self.assertEqual(amount(r.absence['홍길동'][6]), '0.04032')
        self.assertEqual([i.code for i in r.issues], ['W004'])

    def test_T08_cross_year_absence(self):
        r = self.ok(people=[person({'사업1': '0.97312'}, absences=['2026-12-20~2027-01-10'])])
        self.assertEqual(amount(r.absence['홍길동'][0]), '0.02688')
        self.assertEqual([i.code for i in r.issues], ['W003'])

    def test_T09_target_sum(self):
        for value in ('0.90000', '1.10000'):
            with self.subTest(value=value):
                self.bad(people=[person({'사업1': value})], code='E101')

    def test_T10_absence_plus_target(self):
        self.bad([project(year=2026)], [person(absences=['2026-07-15~2026-07-31'])], 2026, 'E101')

    def test_T11_no_project_for_excluded_months(self):
        self.bad([project(excluded=[4, 5, 6])], code='E201')

    def test_T11_no_allowed_months_named_per_project(self):
        # 허용 월이 없는 양수 목표는 연간 합계 오류(E101)와 별개로 사업 단위로 알린다.
        issues = self.bad([project(excluded=list(range(1, 13))), project('사업2')],
                          [person({'사업1': '0.20000', '사업2': '0.70000'})], code='E201')
        self.assertIn('E101', [i.code for i in issues])
        message = next(i.current for i in issues if i.code == 'E201')
        self.assertIn('사업1 목표 0.20000: 허용 월 없음', message)
        self.assertNotIn('사업2', message)
        # 사업 기간 안이어도 전 기간 공백이면 그 사람에게는 허용 월이 없다.
        issues = self.bad([project(period='2027-01-01~2027-03-31'), project('사업2')],
                          [person({'사업1': '0.25000', '사업2': '0.00000'}, absences=['2027-01-01~2027-03-31'])],
                          code='E201')
        self.assertIn('허용 월 없음', next(i.current for i in issues if i.code == 'E201'))

    def test_T12_joint_capacity_conflict(self):
        issues = self.bad([project(excluded=[4, 5, 6]), project('사업2', excluded=[4, 5, 6])],
                         [person({'사업1': '0.50000', '사업2': '0.50000'})], code='E201')
        self.assertIn('0.25000', next(i.current for i in issues if i.code == 'E201'))

    def test_T13_project_limit(self):
        people = [person({'사업1': '0.60000', '사업2': '0.40000'}, name=n) for n in ('가', '나')]
        issues = self.bad([project(limit='1.00000'), project('사업2')], people, code='E103')
        self.assertIn('0.20000', next(i.current for i in issues if i.code == 'E103'))
        self.ok([project(limit='1.50000'), project('사업2')], people)

    def test_T14_zero_or_missing_budget(self):
        for missing in (False, True):
            p = project(limit='0')
            if missing:
                p['yearly'] = {}
            with self.subTest(missing=missing):
                r = self.ok([p, project('사업2')], [person({'사업2': '1.00000'})])
                self.assertEqual([i.code for i in r.issues], ['W005'])
                self.bad([p], code='E102')

    def test_T15_alias_name_reference_integrity(self):
        self.bad([project(), project()], code='E004')
        p = project()
        del p['name']
        self.bad([p], code='E003')
        self.bad(people=[person(), person()], code='E004')
        self.bad(people=[person({'없음': '1'})], code='E007')
        # 사업 파일 오류가 사람 파일에 미등록 약칭 오류를 연달아 만들지 않는다. 다른 사람 오류는 그대로 보인다.
        p = project()
        p['period'] = '2027-13-01~2027-12-31'
        codes = [i.code for i in self.bad([p], people=[person(), person({'사업1': '1.000001'}, name='김철수')])]
        self.assertIn('E006', codes)
        self.assertIn('E005', codes)
        self.assertNotIn('E007', codes)

    def test_T15_issue_order_follows_validation(self):
        # ERROR가 먼저이고, 같은 심각도 안에서는 사업 파일 → 사람 파일, 파일 안은 항목 번호 순서다.
        p = project(excluded=[4, 4])
        p['period'] = '2027-13-01~2027-12-31'
        people = [person({'사업1': '1.000001' if i in (2, 10) else '1'}, name=f'사람{i}') for i in range(11)]
        issues = self.bad([project('사업0', excluded=[5, 5]), p], people)
        ordered = sorted(issues, key=issue_sort_key)
        self.assertEqual([(i.code, Path(i.location.split('.yaml')[0]).name) for i in ordered],
                         [('E006', 'projects'), ('E005', 'people'), ('E005', 'people'),
                          ('W001', 'projects'), ('W001', 'projects')])
        self.assertIn('people[2]', ordered[1].location)
        self.assertIn('people[10]', ordered[2].location)
        self.assertIn('projects[0]', ordered[3].location)

    def test_T15_duplicate_yaml_keys(self):
        paths = self.write([project()], [person()])
        for text in ('year: 2027\nyear: 2028\npeople: []\n',
                     'year: 2027\npeople:\n- name: 홍길동\n  allocations:\n    사업1: "1"\n    사업1: "0"\n'):
            with self.subTest(text=text):
                paths[1].write_text(text, encoding='utf-8')
                inputs, issues = load_inputs(*paths)
                self.assertIsNone(inputs)
                self.assertIn('E002', [i.code for i in issues])

    def test_T16_bad_periods(self):
        for period in ('2027-02-29~2027-03-01', '2027-02-02~2027-02-01', '2027/2/1~2027-03-01'):
            p = project(period=period)
            self.bad([p], code='E006')

    def test_date_formats_preserve_project_and_absence_calculation(self):
        formats = ('{y}-{m:02d}-{d:02d}', '{y}-{m}-{d}',
                   '{y}.{m:02d}.{d:02d}', '{y}.{m}.{d}', '{y}.{m}.{d}.')
        baseline = self.ok([project(year=2026)],
                           [person({'사업1': '0.95430'}, absences=['2026-07-15~2026-07-31'])], 2026)
        # 시작·종료 표기를 독립적으로 바꿔 사업과 휴직 모두에 적용한다.
        for start_format, end_format in itertools.product(formats, repeat=2):
            with self.subTest(start=start_format, end=end_format):
                def period(m1, d1, m2, d2):
                    return (start_format.format(y=2026, m=m1, d=d1) + '~'
                            + end_format.format(y=2026, m=m2, d=d2))
                r = self.ok([project(year=2026, period=period(1, 1, 12, 31))],
                            [person({'사업1': '0.95430'}, absences=[period(7, 15, 7, 31)])], 2026)
                self.assertEqual(r.allocations, baseline.allocations)
                self.assertEqual(r.absence, baseline.absence)

    def test_dotted_leap_day_and_equal_endpoints(self):
        r = self.ok([project(year=2028, period='2028.1.1.~2028.12.31.')],
                    [person({'사업1': '0.99713'}, absences=['2028.2.29.~2028.2.29'])], 2028)
        self.assertEqual(amount(r.absence['홍길동'][1]), '0.00287')

    def test_flexible_dates_still_reject_invalid_calendar_and_syntax(self):
        for value in ('2027.2.29', '2026-4-31', '2026.13.1.', '2026.1.0',
                      '2026.1-1', '2026-1.1', '2026-1-1.', '2026.1.1..',
                      '26.1.1', '2026.001.1', '2026/1/1', '2026. 1. 1',
                      '2026.1.1T00:00:00', '2026.1.1junk', ''):
            with self.subTest(value=value):
                self.bad([project(period=f'{value}~2027-12-31')], code='E006')
                self.bad(people=[person(absences=[f'{value}~2027-12-31'])], code='E006')
        self.bad([project(period='2027.3.1.~2027-2-28')], code='E006')
        self.bad(people=[person(absences=['2027.3.1~2027.2.28.'])], code='E006')
        for value in ('2027.1.1', '2027.1.1~2027.2.1~2027.3.1'):
            self.bad([project(period=value)], code='E006')

    def test_T16_bad_quantities(self):
        for value in ('0.123456', '-1', 'NaN', '.inf', '1e0', True, None, '1_000', '+1', ' 1 '):
            with self.subTest(value=value):
                self.bad(people=[person({'사업1': value})], code='E005')
        self.bad([project(excluded=[13])], code='E003')

    def test_T17_optional_absences_and_zero_months(self):
        p = person()
        del p['absences']
        self.ok([project(excluded=[0, 0, 0])], [p])
        self.ok(people=[person(absences=[])])
        self.ok([project(excluded=[0, 4]), project('사업2')],
                [person({'사업1': '0.60000', '사업2': '0.40000'})])

    def cli(self, *args):
        # Windows 기본 인코딩(cp949)이나 사용자의 PYTHONIOENCODING과 관계없이 자식 출력과 디코딩을 UTF-8로 맞춘다.
        env = {**os.environ, 'PYTHONIOENCODING': 'utf-8'}
        return subprocess.run([sys.executable, str(ROOT / 'src/main.py'), str(self.path / 'people.yaml'), *args],
                              cwd=self.path, capture_output=True, text=True, encoding='utf-8', env=env)

    def test_T18_cli_views(self):
        self.write([project(), project('사업2')], [person(), person({'사업2': '1'}, name='김철수')])
        for args, text in (((), '개인별 연간 요약'), (('--person', '홍길동'), '개인 상세: 홍길동'),
                           (('--project', '사업1'), '사업 상세: 사업1'), (('--all',), '개인 상세: 홍길동')):
            with self.subTest(args=args):
                p = self.cli(*args)
                self.assertEqual(p.returncode, 0, p.stderr + p.stdout)
                self.assertIn(text, p.stdout)
                self.assertIn('1.00000', p.stdout)
                for summary in ('개인별 연간 요약', '사업별 연간 요약'):
                    if not args or args == ('--all',):
                        self.assertIn(summary, p.stdout)
                    else:
                        self.assertNotIn(summary, p.stdout)
                if args and args[0] in ('--person', '--project'):
                    self.assertNotIn('김철수', p.stdout)
                    self.assertNotIn('사업2', p.stdout)
                if args and args[0] == '--project':
                    self.assertIn('연간 배정: 10.00000 / 할당: 1.00000 / 잔여: 9.00000', p.stdout)
                if not args:
                    self.assertNotIn('개인 상세:', p.stdout)
                    self.assertNotIn('사업 상세:', p.stdout)
                elif args[0] == '--person':
                    self.assertNotIn('사업 상세:', p.stdout)
                elif args[0] == '--project':
                    self.assertNotIn('개인 상세:', p.stdout)
        for flag in ('--person', '--project'):
            p = self.cli(flag, '없음')
            self.assertEqual(p.returncode, 1)
            self.assertIn('E008', p.stdout)
            self.assertNotIn('개인별 연간 요약', p.stdout)

    def test_project_detail_shows_participants_and_monthly_totals(self):
        projects = [project(excluded=[4, 5, 6]), project('사업2'), project('미할당사업', limit='0')]
        people = [person({'사업1': '0.60000', '사업2': '0.40000'}, name='가'),
                  person({'사업1': '0.30000', '사업2': '0.70000'}, name='나'),
                  person({'사업2': '1.00000', '사업1': '0'}, name='다')]
        self.write(projects, people)
        result, _ = self.run_files((self.path / 'projects.yaml', self.path / 'people.yaml'))
        self.assertIsNotNone(result)
        assert result is not None
        self.assertFalse(result.has_errors, result.issues)
        for args in (('--project', '사업1'), ('--all-projects',), ('--all',)):
            with self.subTest(args=args):
                output = self.cli(*args)
                self.assertEqual(output.returncode, 0, output.stdout + output.stderr)
                section = output.stdout.split('사업 상세: 사업1 (사업1)', 1)[1].split('-' * 72, 1)[0]
                rows = [line.split(' | ') for line in section.splitlines() if ' | ' in line]
                self.assertEqual([r[0] for r in rows], ['사업/구분', '가', '나', '할당'])
                for row, name in zip(rows[1:3], ('가', '나')):
                    values = result.allocations[name]['사업1']
                    self.assertEqual(row[1:], [amount(v) for v in values] + [amount(sum(values))])
                totals = project_summary(result, '사업1')[3]
                self.assertEqual(rows[-1][1:], [amount(v) for v in totals] + [amount(sum(totals))])
                self.assertEqual(rows[-1][4:7], ['0.00000'] * 3)
        empty = self.cli('--project', '미할당사업')
        self.assertEqual(empty.returncode, 0, empty.stdout + empty.stderr)
        rows = [line.split(' | ') for line in empty.stdout.splitlines() if ' | ' in line]
        self.assertEqual([r[0] for r in rows], ['사업/구분', '할당'])
        self.assertEqual(rows[-1][1:], ['0.00000'] * 13)

    def test_all_displays_every_person_and_project_with_section_breaks(self):
        self.write([project(), project('사업2'), project('미할당사업', limit='0')],
                   [person(), person({'사업2': '1'}, name='김철수')])
        p = self.cli('--all')
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        headings = ['개인별 연간 요약 (사업 / 공백 / 합계)', '사업별 연간 요약 (배정 / 할당 / 잔여)',
                    '개인 상세: 김철수', '개인 상세: 홍길동',
                    '사업 상세: 미할당사업 (미할당사업)', '사업 상세: 사업1 (사업1)', '사업 상세: 사업2 (사업2)']
        positions = []
        lines = p.stdout.splitlines()
        for heading in headings:
            self.assertEqual(lines.count(heading), 1)
            index = lines.index(heading)
            self.assertEqual(lines[index - 2], '')
            self.assertRegex(lines[index - 1], r'^-{3,}$')
            positions.append(index)
        self.assertEqual(positions, sorted(positions))
        self.assertIn('연간 배정: 0.00000 / 할당: 0.00000 / 잔여: 0.00000', p.stdout)
        self.assertIn('[WARNING W005]', p.stdout)

    def test_selected_views_keep_validation_warnings(self):
        self.write([project(excluded=[4, 4]), project('사업2')],
                   [person({'사업1': '0.60000', '사업2': '0.40000'})])
        for args in (('--person', '홍길동'), ('--project', '사업1')):
            with self.subTest(args=args):
                p = self.cli(*args)
                self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
                self.assertIn('[WARNING W001]', p.stdout)
                self.assertIn('ERROR: 0건 / WARNING: 1건', p.stdout)
                self.assertNotIn('개인별 연간 요약', p.stdout)
                self.assertNotIn('사업별 연간 요약', p.stdout)

    def test_category_all_views_show_only_requested_details(self):
        self.write([project(), project('사업2'), project('미할당사업', limit='0')],
                   [person(), person({'사업2': '1'}, name='김철수')])
        for flag, expected, omitted, count in (
                ('--all-people', '개인 상세:', '사업 상세:', 2),
                ('--all-projects', '사업 상세:', '개인 상세:', 3)):
            with self.subTest(flag=flag):
                p = self.cli(flag)
                self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
                self.assertEqual(p.stdout.count(expected), count)
                self.assertNotIn(omitted, p.stdout)
                self.assertNotIn('개인별 연간 요약', p.stdout)
                self.assertNotIn('사업별 연간 요약', p.stdout)
                self.assertIn('[WARNING W005]', p.stdout)
                self.assertIn('12월 | 연간', p.stdout)
                self.assertNotIn('분기', p.stdout)
                if flag == '--all-people':
                    self.assertIn('개인 상세: 김철수', p.stdout)
                    self.assertIn('개인 상세: 홍길동', p.stdout)
                else:
                    self.assertIn('사업 상세: 미할당사업 (미할당사업)', p.stdout)
                    self.assertIn('사업 상세: 사업1 (사업1)', p.stdout)
                    self.assertIn('사업 상세: 사업2 (사업2)', p.stdout)

    def test_category_all_views_do_not_print_invalid_results(self):
        self.write([project()], [person({'사업1': '0.90000'})])
        for flag in ('--all-people', '--all-projects'):
            with self.subTest(flag=flag):
                p = self.cli(flag)
                self.assertEqual(p.returncode, 1)
                self.assertIn('E101', p.stdout)
                self.assertNotIn('개인 상세:', p.stdout)
                self.assertNotIn('사업 상세:', p.stdout)

    def test_T19_warnings(self):
        p = project(excluded=[4, 4])
        p['yearly'][2026] = dict(allocated_manpower='0', excluded_months=[])
        r = self.ok([p, project('사업2')], [person({'사업1': '0.6', '사업2': '0.4'})])
        self.assertEqual(sorted(i.code for i in r.issues), ['W001', 'W002'])

    def test_T20_trim_project_name_alias_and_reference(self):
        p = project('  사업1  ')
        p['name'] = '  연구  지원 사업  '
        r = self.ok([p], [person({' 사업1 ': '1'})])
        self.assertEqual(r.inputs.projects['사업1'].name, '연구  지원 사업')
        self.assertIn('사업1 (연구  지원 사업)', render_result(r))
        self.assertEqual(self.cli('--project', ' 사업1 ').returncode, 0)

    def test_T21_project_whitespace_errors(self):
        for field, value in [('name', ' \t '), ('alias', ' \n '), ('alias', '사업 1'),
                             ('alias', '사업\t1'), ('alias', '사업\n1'), ('alias', '사업\u30001')]:
            with self.subTest(field=field, value=value):
                p = project()
                p[field] = value
                self.bad([p])

    def test_T22_normalized_duplicates(self):
        self.bad([project(), project(' 사업1 ')], code='E004')
        self.bad(people=[person({'사업1': '0.6', ' 사업1 ': '0.4'})], code='E004')

    def test_T23_year_types_and_limits(self):
        for year in (999, 10000, True, '2027', 2027.0):
            with self.subTest(year=year):
                self.bad(year=year, code='E003')
        for year in (1000, 9999):
            self.ok([project(year=year)], [person()], year)
        self.bad([project(excluded=[True])], code='E003')
        p = project()
        p['yearly'] = {True: dict(allocated_manpower='1', excluded_months=[])}
        self.bad([p], code='E003')

    def test_T23_required_fields_and_nonempty_lists(self):
        self.bad(projects=[], code='E003')
        self.bad(people=[], code='E003')
        p = project()
        del p['yearly']
        self.bad([p], code='E003')

    def test_T24_full_absence_empty_targets(self):
        p = project()
        p['yearly'] = {}
        r = self.ok([p], [person({}, absences=['2027-01-01~2027-12-31'])])
        self.assertEqual(r.allocations['홍길동'], {})
        self.assertEqual(r.absence['홍길동'], list(BASE))
        self.bad(people=[person({})], code='E101')

    def test_T25_options_and_monthly_annual_output(self):
        self.write([project()], [person()])
        options = (('--all',), ('--person', '홍길동'), ('--project', '사업1'),
                   ('--all-people',), ('--all-projects',))
        for first, second in itertools.combinations(options, 2):
            with self.subTest(first=first, second=second):
                self.assertEqual(self.cli(*first, *second).returncode, 2)
        for args in options:
            with self.subTest(args=args):
                p = self.cli(*args)
                self.assertEqual(p.returncode, 0, p.stderr)
                self.assertNotIn('분기', p.stdout)
                lines = [line.split(' | ') for line in p.stdout.splitlines() if ' | ' in line]
                self.assertEqual(lines[0], ['사업/구분', '1월', '2월', '3월', '4월', '5월', '6월',
                                            '7월', '8월', '9월', '10월', '11월', '12월', '연간'])
                for row in lines[1:]:
                    self.assertEqual(len(row), 14)
                    if row[0] == '사업/구분':
                        self.assertEqual(row, lines[0])
                        continue
                    self.assertEqual(row[-1], '1.00000' if row[0] != '공백' else '0.00000')
        (self.path / 'projects.yaml').unlink()
        p = self.cli()
        self.assertEqual(p.returncode, 1)
        self.assertIn('E001', p.stdout)

    def test_T26_no_duplicate_warning_for_budget_error(self):
        for missing in (False, True):
            p = project(limit='0')
            if missing:
                p['yearly'] = {}
            issues = self.bad([p], code='E102')
            self.assertEqual(sum(i.code == 'E102' for i in issues), 1)
            self.assertNotIn('W005', [i.code for i in issues])

    def test_yaml_numeric_literals_keep_decimal_precision(self):
        paths = self.write([project()], [person()])
        paths[1].write_text('year: 2027\npeople:\n- name: 홍길동\n  allocations:\n    사업1: 1.00000\n', encoding='utf-8')
        r, issues = self.run_files(paths)
        self.assertIsNotNone(r, issues)
        self.assertFalse(r.has_errors)
        for literal in ('1.000001', '1e0', '.nan', '.inf', '-1', '0x01'):
            paths[1].write_text(f'year: 2027\npeople:\n- name: 홍길동\n  allocations:\n    사업1: {literal}\n', encoding='utf-8')
            inputs, issues = load_inputs(*paths)
            self.assertIsNone(inputs, literal)
            self.assertIn('E005', [i.code for i in issues])

    def test_T27_input_encodings(self):
        paths = self.write([project('연구사업'), project('운영사업')],
                           [person({'연구사업': '0.60000', '운영사업': '0.40000'}, name='홍길동')])
        texts = [p.read_text(encoding='utf-8') for p in paths]
        expected, _ = self.run_files(paths)
        for encodings in (('cp949', 'cp949'), ('utf-8-sig', 'utf-8-sig'), ('cp949', 'utf-8')):
            with self.subTest(encodings=encodings):
                for path, text, encoding in zip(paths, texts, encodings):
                    path.write_text(text, encoding=encoding)
                result, issues = self.run_files(paths)
                self.assertIsNotNone(result, issues)
                self.assertFalse(result.has_errors, result.issues)
                self.assertEqual(result.allocations, expected.allocations)
        # UTF-8과 CP949 어느 쪽으로도 해석할 수 없는 바이트는 E001이다.
        paths[1].write_bytes(texts[1].encode('utf-8') + b'# \xff\n')
        inputs, issues = load_inputs(*paths)
        self.assertIsNone(inputs)
        self.assertIn('E001', [i.code for i in issues])

    def test_T30_editor_document_round_trip(self):
        # GUI 편집기 형식으로 읽고 다시 쓴 YAML이 원본과 같은 입력·결과·진단을 만든다. 저장소 예제는 읽기만 한다.
        from allocation.document import dump_people, dump_projects, read_people, read_projects
        from allocation.loading import decode_text
        folders = [ROOT / 'myteam', *sorted((ROOT / 'templates').glob('[0-9]*')),
                   *sorted((ROOT / 'templates' / 'errors').glob('[0-9]*'))]
        checked = 0
        for folder in folders:
            p = folder / 'projects.yaml'
            for h in sorted((folder / 'people').glob('*.yaml')):
                with self.subTest(people=str(h.relative_to(ROOT))):
                    projects, dropped = read_projects(decode_text(p.read_bytes()))
                    doc, dropped_people = read_people(decode_text(h.read_bytes()))
                    texts = {'projects': dump_projects(projects), 'people': dump_people(doc)}
                    a, a_issues = load_inputs(p, h)
                    b, b_issues = load_inputs(p, h, texts)
                    if dropped or dropped_people:
                        # 중복 키·미지원 필드는 편집기가 버리고 사용자에게 알린다.
                        self.assertIn(folder.name, ('15_duplicate_yaml_key', '16_invalid_fields'))
                        continue
                    self.assertEqual(a, b)
                    if a is None:
                        self.assertEqual(sorted(i.code for i in a_issues), sorted(i.code for i in b_issues))
                        continue
                    ra, rb = calculate(a, a_issues), calculate(b, b_issues)
                    self.assertEqual((ra.allocations, ra.issues), (rb.allocations, rb.issues))
                    checked += 1
        self.assertGreaterEqual(checked, 15)

    def test_T30_editor_document_values(self):
        from allocation.document import (PeopleDoc, PersonEntry, ProjectEntry, YearEntry, dump_people,
                                         dump_projects, person_totals, quantity_units, read_people, read_projects)
        # 따옴표·콜론·# 이 든 이름도 안전하게 쓰고, 숫자처럼 보이는 값은 문자열로 남겨 원문 그대로 검증받는다.
        projects = [ProjectEntry('연구: "1차" #A', 'A1', '2026.1.1.', '2026-12-31', [YearEntry('2026', '1.5', '4 5,6')])]
        text = dump_projects(projects)
        again, dropped = read_projects(text)
        self.assertEqual((again, dropped), ([ProjectEntry('연구: "1차" #A', 'A1', '2026.1.1.', '2026-12-31',
                                                          [YearEntry('2026', '1.5', '4, 5, 6')])], []))
        doc = PeopleDoc('2026', [PersonEntry('김 철수', {'A1': '1', 'B': ''})])
        self.assertNotIn('"B"', dump_people(doc))
        self.assertEqual(read_people(dump_people(doc))[0].people[0].allocations, {'A1': '1'})
        inputs, issues = load_inputs('p.yaml', 'h.yaml', {'projects': text, 'people': dump_people(doc)})
        self.assertIsNotNone(inputs, issues)
        self.assertEqual(inputs.projects['A1'].yearly[2026].excluded, frozenset({4, 5, 6}))
        # 잘못된 연도·월은 문자열로 써서 파일 입력과 같은 오류 코드를 받는다.
        bad = [ProjectEntry('X', 'X', '2026-01-01', '2026-12-31', [YearEntry('20x6', '1', '4월')])]
        _, issues = load_inputs('p.yaml', 'h.yaml', {'projects': dump_projects(bad), 'people': dump_people(doc)})
        self.assertIn('E003', [i.code for i in issues])
        self.assertEqual([quantity_units(v) for v in ('1', '0.5', '0.00001', '1e0', '-1', '0.123456', '')],
                         [100000, 50000, 1, None, None, None, None])
        self.assertEqual(person_totals(PersonEntry('a', {'A1': '0.6', 'B': '0.4'}), '2026'), (100000, 0))
        self.assertEqual(person_totals(PersonEntry('a', {'A1': 'x'}), '2026'), (None, 0))
        self.assertEqual(person_totals(PersonEntry('a', {'A1': '1'}), 'abc'), (100000, None))

    def test_T28_excel_output(self):
        # 사업1의 제외 월 4월 중복으로 W001 경고 1건이 생긴다.
        paths = self.write([project(excluded=[4, 4]), project('사업2')],
                           [person({'사업1': '0.60000', '사업2': '0.40000'}), person({'사업2': '1'}, name='김철수')])
        result, _ = self.run_files(paths)
        p = self.cli('--person', '홍길동', '--excel', 'out.xlsx')
        self.assertEqual(p.returncode, 0, p.stderr + p.stdout)
        self.assertNotIn('개인별 연간 요약', p.stdout)
        self.assertTrue(p.stdout.rstrip().endswith('엑셀 저장: out.xlsx'))
        book = load_workbook(self.path / 'out.xlsx')
        self.assertEqual(book.sheetnames, ['검증', '요약', '개인상세', '사업상세'])

        def cells(sheet):
            return [['' if c.value is None else f'{c.value:.5f}' if type(c.value) in (int, float) else c.value
                     for c in row] for row in sheet.iter_rows()]

        def padded(rows, width):
            return [row + [''] * (width - len(row)) for row in rows]

        check = cells(book['검증'])
        self.assertEqual(check[3][:2], ['검증 결과', 'ERROR 0건 / WARNING 1건'])
        self.assertEqual(check[6][:2], ['WARNING', 'W001'])
        self.assertEqual(len(check), 7)
        # 조회 옵션이 --person이어도 엑셀에는 모든 사람과 사업이 들어간다.
        expected = [['이름', '사업', '공백', '합계']]
        expected += [[name, *map(amount, person_summary(result, name))] for name in result.allocations]
        expected += [[], ['사업약칭', '사업명', '배정', '할당', '잔여', '비고']]
        details = [['사업약칭', '사업명', '배정', '할당', '잔여', '비고', *(f'{m}월' for m in range(1, 13))]]
        for alias in result.inputs.projects:
            item, limit, missing, totals = project_summary(result, alias)
            row = [alias, item.name, *map(amount, (limit, sum(totals), limit - sum(totals))),
                   '배정 미등록' if missing else '']
            expected.append(row)
            details.append(row + [amount(v) for v in totals])
        self.assertEqual(cells(book['요약']), padded(expected, 6))
        self.assertEqual(cells(book['사업상세']), details)
        self.assertEqual(details[1][:5], ['사업1', '사업1', '10.00000', '0.60000', '9.40000'])
        people = []
        for name in result.inputs.people:
            people += [[name], ['사업/구분', *(f'{m}월' for m in range(1, 13)), '연간']]
            people += [[label, *(amount(v) for v in values + [sum(values)])]
                       for label, values in person_table(result, name)]
            people.append([])
        # 마지막 빈 행은 값이 없으므로 파일에 남지 않는다.
        self.assertEqual(cells(book['개인상세']), padded(people[:-1], 14))
        for sheet in book.worksheets[1:]:
            for row in sheet.iter_rows():
                for c in row:
                    if type(c.value) in (int, float):
                        self.assertEqual(c.number_format, '0.00000', (sheet.title, c.coordinate))
        # 같은 입력은 같은 셀 값을 만든다.
        self.assertEqual(self.cli('--excel', 'again.xlsx').returncode, 0)
        again = load_workbook(self.path / 'again.xlsx')
        self.assertEqual([cells(s) for s in again.worksheets], [cells(s) for s in book.worksheets])

        # 확장자 오류는 사용법 오류, 저장 실패는 실행 오류다.
        self.assertEqual(self.cli('--excel', 'out.csv').returncode, 2)
        (self.path / 'folder.xlsx').mkdir()
        p = self.cli('--excel', 'folder.xlsx')
        self.assertEqual(p.returncode, 1)
        self.assertIn('엑셀 저장 실패: folder.xlsx', p.stdout)
        self.assertEqual(sorted(f.name for f in self.path.glob('.~*')), [])

        # 쓰는 도중 실패해도 기존 파일은 그대로이고 임시 파일도 남지 않는다.
        before = (self.path / 'out.xlsx').read_bytes()

        def broken_save(book, target):
            Path(target).write_bytes(b'partial')
            raise OSError('디스크 가득 참')

        with patch('openpyxl.Workbook.save', broken_save), self.assertRaises(OSError):
            write_workbook(result, self.path / 'out.xlsx')
        self.assertEqual((self.path / 'out.xlsx').read_bytes(), before)
        self.assertEqual(sorted(f.name for f in self.path.glob('.~*')), [])
        write_workbook(result, self.path / 'new.xlsx')
        if os.name == 'posix':
            mask = os.umask(0)
            os.umask(mask)
            self.assertEqual((self.path / 'new.xlsx').stat().st_mode & 0o777, 0o666 & ~mask)

        # ERROR가 있으면 파일을 만들지 않는다.
        self.write([project()], [person({'사업1': '0.90000'})])
        for broken in (False, True):
            if broken:
                (self.path / 'projects.yaml').unlink()
            p = self.cli('--excel', 'error.xlsx')
            self.assertEqual(p.returncode, 1)
            self.assertIn('엑셀: 오류로 생성하지 않음', p.stdout)
            self.assertFalse((self.path / 'error.xlsx').exists())

    def test_schema_errors_aggregate(self):
        p = project()
        p['id'] = 1
        p['period'] = '잘못된 날짜'
        del p['yearly'][2027]['allocated_manpower']
        issues = self.bad([p])
        self.assertGreaterEqual(len(issues), 3)

    def test_omitted_excluded_months_matches_empty_list(self):
        explicit = self.ok()
        p = project()
        del p['yearly'][2027]['excluded_months']
        omitted = self.ok([p])
        self.assertEqual(omitted.issues, [])
        self.assertEqual(omitted.inputs.projects['사업1'].yearly[2027].excluded, frozenset())
        self.assertEqual(omitted.allocations, explicit.allocations)
        self.assertEqual(omitted.allocations['홍길동']['사업1'], list(BASE))

    def test_present_excluded_months_still_requires_list(self):
        for value in (None, 4, '4', {}):
            with self.subTest(value=value):
                p = project()
                p['yearly'][2027]['excluded_months'] = value
                self.bad([p], code='E003')

    def test_duplicate_year_values_never_overwrite(self):
        paths = self.write([project()], [person()])
        paths[0].write_text('''projects:
- name: 사업1
  alias: 사업1
  period: "2027-01-01~2027-12-31"
  yearly:
    2027: {allocated_manpower: "1", excluded_months: []}
    02027: {allocated_manpower: "2", excluded_months: []}
''', encoding='utf-8')
        inputs, issues = load_inputs(*paths)
        self.assertIsNone(inputs)
        self.assertIn('E002', [i.code for i in issues])

    def test_yaml_syntax_encoding_and_unsafe_tags(self):
        paths = self.write([project()], [person()])
        for text in ('[unclosed', '!!python/object:os.system {}', '---\n{}\n---\n{}', '? [a, b]\n: c', 'year: &x [*x]\npeople: []'):
            paths[1].write_text(text, encoding='utf-8')
            inputs, issues = load_inputs(*paths)
            self.assertIsNone(inputs, text)
        paths[1].write_bytes(b'\xff')
        self.assertIsNone(load_inputs(*paths)[0])

    def test_out_of_year_and_nested_absence(self):
        r = self.ok(people=[person(absences=['2025-01-01~2025-12-31'])])
        self.assertEqual(r.absence['홍길동'], [0] * 12)
        r = self.ok(people=[person({'사업1': '0.91667'}, absences=[
            '2027-01-01~2027-01-31', '2027-01-05~2027-01-10', '2027-01-10~2027-01-20'])])
        self.assertEqual(r.absence['홍길동'][0], 8333)

    def test_post_validation_rejects_corrupt_result(self):
        r = self.ok()
        r.allocations['홍길동']['사업1'][0] -= 1
        self.assertTrue(validate_result(r))

    def test_offline_execution_without_socket(self):
        self.write([project()], [person()])
        previous = Path.cwd()
        try:
            os.chdir(self.path)
            with patch('socket.socket', side_effect=AssertionError('외부 통신 불가')), redirect_stdout(io.StringIO()) as output:
                self.assertEqual(main(['people.yaml', '--all']), 0)
            self.assertIn('1.00000', output.getvalue())
        finally:
            os.chdir(previous)


class TransportTests(unittest.TestCase):
    def test_reassignment_needed(self):
        # A를 1월에 우선 확정하면 B가 부족하다. A를 2월로 옮겨야 한다.
        values, shortage, _, _ = transport({'A': 1, 'B': 3}, {'A': {0, 1}, 'B': {0, 2}}, [2, 1, 1])
        self.assertEqual(shortage, 0)
        self.assertEqual(values, {'A': [0, 1, 0], 'B': [2, 0, 1]})

    def test_small_problems_against_exhaustive_oracle(self):
        rng = random.Random(527)
        for case in range(200):
            capacities = [rng.randrange(3) for _ in range(3)]
            total = sum(capacities)
            cuts = sorted([0, total, rng.randrange(total + 1), rng.randrange(total + 1)])
            targets = {chr(65 + p): cuts[p + 1] - cuts[p] for p in range(3)}
            allowed = {alias: {m for m in range(3) if rng.randrange(2)} for alias in targets}
            order = sorted(targets, key=lambda a: (len(allowed[a]), a))
            solutions = []
            def enumerate_rows(index, remaining, rows):
                if index == len(order):
                    if not any(remaining):
                        solutions.append(copy.deepcopy(rows))
                    return
                alias = order[index]
                for row in itertools.product(*(range(c + 1) if m in allowed[alias] else [0]
                                               for m, c in enumerate(remaining))):
                    if sum(row) == targets[alias]:
                        rows[alias] = list(row)
                        enumerate_rows(index + 1, [c - v for c, v in zip(remaining, row)], rows)
            enumerate_rows(0, capacities, {})
            values, shortage, _, _ = transport(targets, allowed, capacities)
            with self.subTest(case=case):
                self.assertEqual(values is not None, bool(solutions))
                if solutions:
                    def deviation(rows):
                        errors = []
                        for alias in order:
                            months = sorted(allowed[alias])
                            if months:
                                q, r = divmod(targets[alias], len(months))
                                errors.extend(abs(rows[alias][m] - q - (j < r)) for j, m in enumerate(months))
                        return tuple(errors)
                    self.assertEqual(deviation(values), min(map(deviation, solutions)))
                    self.assertEqual(shortage, 0)
                else:
                    self.assertGreater(shortage, 0)


if __name__ == '__main__':
    unittest.main()
