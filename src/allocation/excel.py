"""계산 결과를 주제별 4개 시트의 .xlsx 파일로 저장한다. 조회 옵션과 관계없이 전체 결과를 담는다."""
import os
import tempfile
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .rendering import issue_counts
from .tables import (issue_table, people_summary_table, person_detail_table, project_detail_table,
                     project_summary_table)

NUMBER_FORMAT = '0.00000'
BOLD = Font(bold=True)


def cell_value(value):
    # 0.00001 단위 정수를 float 없이 Decimal로 바꾼다. 셀에는 정확한 십진 문자열로 기록된다.
    if type(value) is int:
        return Decimal(value).scaleb(-5)
    return value if value != '' else None


class Sheet:
    def __init__(self, sheet):
        self.sheet = sheet

    def row(self, values, bold=False):
        self.sheet.append(values)
        for cell in self.sheet[self.sheet.max_row]:
            if isinstance(cell.value, Decimal):
                cell.number_format = NUMBER_FORMAT
            if bold:
                cell.font = BOLD

    def table(self, headers, rows, bold_label=None):
        self.row(headers, bold=True)
        for row in rows:
            self.row([cell_value(value) for value in row], bold=row[0] == bold_label)

    def blank(self):
        self.sheet.append([None])

    def widths(self, *widths):
        for index, width in enumerate(widths, 1):
            self.sheet.column_dimensions[get_column_letter(index)].width = width


def write_workbook(result, path):
    inputs = result.inputs
    book = Workbook()

    check = Sheet(book.active)
    check.sheet.title = '검증'
    errors, warnings = issue_counts(result.issues)
    for label, value in (('관리 연도', inputs.year), ('사업 파일', inputs.project_file),
                         ('사람 파일', inputs.people_file), ('검증 결과', f'ERROR {errors}건 / WARNING {warnings}건')):
        check.row([label, value])
    check.blank()
    check.table(*issue_table(result.issues))
    check.widths(12, 8, 30, 40, 30, 40)

    summary = Sheet(book.create_sheet('요약'))
    summary.table(*people_summary_table(result))
    summary.blank()
    summary.table(*project_summary_table(result))
    summary.sheet.freeze_panes = 'B2'
    summary.widths(16, 32, 10, 10, 10, 12)

    people = Sheet(book.create_sheet('개인상세'))
    for name in inputs.people:
        people.row([name], bold=True)
        people.table(*person_detail_table(result, name), bold_label='합계')
        people.blank()
    people.widths(16, *[10] * 13)

    projects = Sheet(book.create_sheet('사업상세'))
    projects.table(*project_detail_table(result))
    projects.sheet.freeze_panes = 'C2'
    projects.widths(16, 32, 10, 10, 10, 12, *[10] * 12)

    save_replacing(book, path)


def save_replacing(book, path):
    """같은 폴더의 임시 파일에 다 쓴 뒤 바꿔 넣는다. 저장에 실패하면 기존 파일을 그대로 두고 임시 파일을 지운다."""
    target = Path(path)
    handle, temp = tempfile.mkstemp(prefix='.~', suffix='.xlsx', dir=target.parent)
    os.close(handle)
    try:
        # mkstemp는 소유자만 읽을 수 있게 만든다. 직접 만든 파일처럼 umask를 따른다.
        mask = os.umask(0)
        os.umask(mask)
        os.chmod(temp, 0o666 & ~mask)
        book.save(temp)
        os.replace(temp, target)
    except BaseException:
        Path(temp).unlink(missing_ok=True)
        raise
